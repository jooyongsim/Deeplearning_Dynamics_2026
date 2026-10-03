"""Build the editable PPTX from the HTML deck's exported layout.

    bash render_deck.sh export     # HTML -> build/layout.json
    python build_pptx.py           # -> slides/ADL_2026_Dynamics.pptx

Everything is native PowerPoint: text boxes, shapes, pictures, and equations as
Office Math (LaTeX -> OMML via pandoc, see tex2omml.py). The slide chrome
(title bar, mark, footer) follows the AMSL template ADL_2026_Lec01_AMSL.pptx,
which is also the base file so its theme is kept.
"""
import copy
import io
import json
import os
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu

from tex2omml import convert

sys.stdout.reconfigure(encoding="utf-8")
HERE = os.path.dirname(os.path.abspath(__file__))
LECTURE = os.path.dirname(HERE)
SLIDES_DIR = os.path.join(LECTURE, "slides")
TEMPLATE = r"C:\Users\Sim\Dropbox\[Courses]\[고급딥러닝시스템과응용]\08_Claude\05_design\ADL_2026_Lec01_AMSL.pptx"
OUT = os.path.join(SLIDES_DIR, "ADL_2026_Dynamics.pptx")
CREDIT = "Sookmyung Women’s Univ. Autonomous Mechanical Systems Lab"

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
NS_A14 = "http://schemas.microsoft.com/office/drawing/2010/main"
NS_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

MATH_SCALE = 0.97   # Cambria Math has a smaller x-height than Arial; KaTeX is already 1.1em
FONTS = {"core": "Arial", "display": "Pretendard", "hangul": "Noto Sans KR", "mono": "Consolas"}
SUP = str.maketrans("0123456789", "⁰¹²³⁴⁵⁶⁷⁸⁹")
SUB = str.maketrans("0123456789", "₀₁₂₃₄₅₆₇₈₉")
VULGAR = {("1", "2"): "½", ("1", "3"): "⅓", ("2", "3"): "⅔", ("1", "4"): "¼", ("3", "4"): "¾"}


def _compact_fractions(omml):
    """pandoc writes \\tfrac as a linear fraction (type=lin). Keep it on the text line:
    common numeric ones become single glyphs (½ ¾ ⅓), the rest stay linear a/b."""
    om = etree.fromstring(omml)
    for f in list(om.iter(f"{{{NS_M}}}f")):
        typ = f.find(f"{{{NS_M}}}fPr/{{{NS_M}}}type")
        if typ is None or typ.get(f"{{{NS_M}}}val") != "lin":
            continue
        num = "".join(f.find(f"{{{NS_M}}}num").itertext()).strip()
        den = "".join(f.find(f"{{{NS_M}}}den").itertext()).strip()
        glyph = VULGAR.get((num, den))
        if not glyph and (num + den).isdigit():   # e.g. 1/12 -> ¹⁄₁₂, same look as ½
            glyph = num.translate(SUP) + "⁄" + den.translate(SUB)
        if glyph:
            r = etree.Element(f"{{{NS_M}}}r")
            rp = etree.SubElement(r, f"{{{NS_M}}}rPr")
            etree.SubElement(rp, f"{{{NS_M}}}nor")   # normal text: no operator spacing around ⁄
            etree.SubElement(r, f"{{{NS_M}}}t").text = glyph
            f.getparent().replace(f, r)
    return etree.tostring(om, encoding="unicode")


def convert_norm(texs):
    res = convert(texs)
    return {t: _compact_fractions(res[t]) for t in texs}


def E(px):
    return Emu(int(round(px * 9525)))


# ---------------------------------------------------------------- template
def load_template():
    prs = Presentation(TEMPLATE)
    cover_mark = prs.slides[0].shapes[2].image.blob
    slide_mark = prs.slides[1].shapes[2].image.blob
    # drop the template's own slides; keep master/layout/theme
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        prs.part.drop_rel(sldId.rId)
        sldIdLst.remove(sldId)
    return prs, cover_mark, slide_mark


# ---------------------------------------------------------------- shapes
def _fill(shape, hexcol, alpha=None):
    if hexcol is None:
        shape.fill.background()
        return
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(hexcol)
    if alpha is not None:
        clr = shape.fill._xPr.find(qn("a:solidFill"))[0]
        etree.SubElement(clr, qn("a:alpha")).set("val", str(int(alpha * 100000)))


def _line(shape, hexcol, sw):
    if not hexcol or not sw:
        shape.line.fill.background()
        return
    shape.line.color.rgb = RGBColor.from_string(hexcol)
    shape.line.width = E(sw)


def add_rect(slide, it):
    radius = it.get("radius") or 0
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sw = it.get("sw") or 0
    # CSS border sits inside the box; PowerPoint strokes are centred on the edge
    x, y, w, h = it["x"] + sw / 2, it["y"] + sw / 2, it["w"] - sw, it["h"] - sw
    shp = slide.shapes.add_shape(kind, E(x), E(y), E(max(w, 0.5)), E(max(h, 0.5)))
    if radius:
        shp.adjustments[0] = min(0.5, radius / min(w, h))
    _fill(shp, it.get("fill"))
    _line(shp, it.get("stroke"), sw)
    shp.shadow.inherit = False
    return shp


def add_oval(slide, it):
    sw = it.get("sw") or 0
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, E(it["x"] + sw / 2), E(it["y"] + sw / 2),
                                 E(it["w"] - sw), E(it["h"] - sw))
    _fill(shp, it.get("fill"))
    _line(shp, it.get("stroke"), sw)
    shp.shadow.inherit = False
    return shp


def add_tri(slide, it):
    pts = it["pts"]
    fb = slide.shapes.build_freeform(E(pts[0][0]), E(pts[0][1]), scale=1.0)
    fb.add_line_segments([(E(x), E(y)) for x, y in pts[1:]], close=True)
    shp = fb.convert_to_shape()
    _fill(shp, it.get("fill"))
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_polyline(slide, it):
    pts = it["pts"]
    if len(pts) == 2:
        (x1, y1), (x2, y2) = pts
        shp = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
    else:
        fb = slide.shapes.build_freeform(E(pts[0][0]), E(pts[0][1]), scale=1.0)
        fb.add_line_segments([(E(x), E(y)) for x, y in pts[1:]], close=False)
        shp = fb.convert_to_shape()
        shp.fill.background()
    shp.line.color.rgb = RGBColor.from_string(it["color"] or "111111")
    shp.line.width = E(it["sw"])
    ln = shp.line._get_or_add_ln()
    if it.get("dash"):
        etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
    if it.get("head"):
        t = etree.SubElement(ln, qn("a:tailEnd"))
        t.set("type", "triangle"); t.set("w", "med"); t.set("len", "med")
    shp.shadow.inherit = False
    return shp


def add_picture(slide, blob_or_path, x, y, w, h, alpha=None):
    src = io.BytesIO(blob_or_path) if isinstance(blob_or_path, bytes) else blob_or_path
    pic = slide.shapes.add_picture(src, E(x), E(y), E(w), E(h))
    if alpha is not None:
        blip = pic._element.find(".//" + qn("a:blip"))
        etree.SubElement(blip, qn("a:alphaModFix")).set("amt", str(int(alpha * 100000)))
    return pic


# ---------------------------------------------------------------- text
def _rpr(tag, run, scale=1.0):
    """a:rPr / a:endParaRPr for a run dict."""
    el = etree.Element(qn(tag))
    el.set("lang", "en-US")
    el.set("sz", str(int(round(run.get("size", 24) * 0.75 * scale * 100))))
    if run.get("bold"):
        el.set("b", "1")
    if run.get("italic"):
        el.set("i", "1")
    el.set("dirty", "0")
    col = run.get("color") or "000000"
    sf = etree.SubElement(el, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr")).set("val", col)
    face = FONTS.get(run.get("font", "core"), "Arial")
    etree.SubElement(el, qn("a:latin")).set("typeface", face)
    etree.SubElement(el, qn("a:ea")).set("typeface", "Malgun Gothic")
    etree.SubElement(el, qn("a:cs")).set("typeface", face)
    return el


def _math_el(run, omml, para):
    """a14:m wrapping m:oMath (inline) or m:oMathPara (display)."""
    om = etree.fromstring(omml)
    # strip Word-only formatting; give every math run PowerPoint run properties
    for w in om.xpath(".//*[namespace-uri()=$ns]", ns=NS_W):
        w.getparent().remove(w)
    for mr in om.iter(f"{{{NS_M}}}r"):
        rpr = _rpr("a:rPr", dict(run, bold=False, italic=False), MATH_SCALE)
        for child in list(rpr):
            if child.tag in (qn("a:latin"), qn("a:cs"), qn("a:ea")):
                rpr.remove(child)
        etree.SubElement(rpr, qn("a:latin")).set("typeface", "Cambria Math")
        mt = mr.find(f"{{{NS_M}}}t")
        mt.addprevious(rpr)
    m = etree.Element(f"{{{NS_A14}}}m", nsmap={"a14": NS_A14})
    if para:
        mp = etree.SubElement(m, f"{{{NS_M}}}oMathPara", nsmap={"m": NS_M})
        pr = etree.SubElement(mp, f"{{{NS_M}}}oMathParaPr")
        etree.SubElement(pr, f"{{{NS_M}}}jc").set(f"{{{NS_M}}}val", "left")
        mp.append(om)
    else:
        m.append(om)
    return m


def add_text(slide, it, omml, override=None):
    runs = it["runs"]
    lines = [[]]
    for r in runs:
        if r.get("br"):
            lines.append([])
        else:
            lines[-1].append(r)
    lines = [ln for ln in lines if ln]
    has_math = any("m" in r for r in runs)
    lh = it.get("lh") or 0
    # one rendered line in the browser -> never let PowerPoint wrap it
    if "nlines" in it:   # every browser line is its own paragraph -> nothing wrapped in the HTML
        single = it["nlines"] <= len(lines)
    else:
        single = len(lines) == 1 and it["h"] <= 1.7 * max(lh, 1)
    x, y, w, h = it["x"], it["y"], it["w"], it["h"]
    # PowerPoint sets these glyphs higher in their line box than CSS does
    y += {"op-oval": 7, "op-txt": 4, "eqlabel": 5}.get(it.get("cls"), 0)
    if not single:
        w = min(w + 2, 1257 - x)   # never past the right margin
    tb = slide.shapes.add_textbox(E(x), E(y), E(max(w, 4)), E(max(h, 4)))
    tf = tb.text_frame
    body = tf._txBody
    bp = body.find(qn("a:bodyPr"))
    for k in ("lIns", "tIns", "rIns", "bIns"):
        bp.set(k, "0")
    bp.set("wrap", "none" if single else "square")
    valign = it.get("valign")
    if single and lh and h > 1.5 * lh:
        valign = "middle"
    bp.set("anchor", {"middle": "ctr", "bottom": "b"}.get(valign, "t"))
    for p in body.findall(qn("a:p")):
        body.remove(p)
    algn = {"center": "ctr", "right": "r", "justify": "just"}.get(it.get("align"), "l")
    for ln in lines:
        p = etree.SubElement(body, qn("a:p"))
        ppr = etree.SubElement(p, qn("a:pPr"))
        ppr.set("algn", algn)
        ppr.set("marL", "0"); ppr.set("indent", "0")
        line_has_math = any("m" in r for r in ln)
        spc = etree.SubElement(ppr, qn("a:lnSpc"))
        if line_has_math or not lh:
            etree.SubElement(spc, qn("a:spcPct")).set("val", "100000")
        else:
            etree.SubElement(spc, qn("a:spcPts")).set("val", str(int(round(lh * 0.75 * 100))))
        etree.SubElement(ppr, qn("a:buNone"))
        only_math = len(ln) == 1 and "m" in ln[0]
        for r in ln:
            if "m" in r:
                p.append(_math_el(r, omml[r["m"]], para=only_math))
            else:
                ar = etree.SubElement(p, qn("a:r"))
                ar.append(_rpr("a:rPr", r))
                etree.SubElement(ar, qn("a:t")).text = r["t"]
        p.append(_rpr("a:endParaRPr", ln[-1] if "m" not in ln[-1] else dict(ln[-1], color=None)))
    if has_math:
        _wrap_alternate(tb._element)
    return tb


def _wrap_alternate(sp):
    """PowerPoint only reads a14:m inside mc:AlternateContent/mc:Choice Requires="a14"."""
    parent = sp.getparent()
    ac = etree.Element(f"{{{NS_MC}}}AlternateContent", nsmap={"mc": NS_MC})
    ch = etree.SubElement(ac, f"{{{NS_MC}}}Choice", nsmap={"a14": NS_A14})
    ch.set("Requires", "a14")
    parent.replace(sp, ac)
    ch.append(sp)


def plain_text(slide, x, y, w, h, text, size_pt, color, bold=False, italic=False, align="l",
               anchor="ctr", face="Arial", spc=None, alpha=None):
    tb = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    bp = tb.text_frame._txBody.find(qn("a:bodyPr"))
    for k in ("lIns", "tIns", "rIns", "bIns"):
        bp.set(k, "25400")
    bp.set("wrap", "square"); bp.set("anchor", anchor)
    p = tb.text_frame.paragraphs[0]._p
    ppr = etree.SubElement(p, qn("a:pPr")); ppr.set("algn", align)
    r = etree.SubElement(p, qn("a:r"))
    rpr = _rpr("a:rPr", {"size": size_pt / 0.75, "bold": bold, "italic": italic, "color": color})
    for c in list(rpr):
        if c.tag in (qn("a:latin"), qn("a:cs")):
            c.set("typeface", face)
    if spc is not None:
        rpr.set("spc", str(spc))
    r.append(rpr)
    etree.SubElement(r, qn("a:t")).text = text
    return tb


# ---------------------------------------------------------------- chrome
def chrome_content(slide, n, title, mark):
    add_rect(slide, {"x": 20.9, "y": 13.5, "w": 9.8, "h": 63.4, "fill": "203864"})
    plain_text(slide, 46, 13.5, 1146, 63.4, title, 28, "203864", bold=True, spc=-28)
    add_picture(slide, mark, 1198.7, 14.4, 62, 62, alpha=0.5)
    add_rect(slide, {"x": 0, "y": 694.2, "w": 1280, "h": 26.5, "fill": "203864"})
    add_rect(slide, {"x": 948.4, "y": 694.2, "w": 331.6, "h": 26.5, "fill": "B4C3D6"})
    plain_text(slide, 8, 694.2, 520.8, 26.5, CREDIT, 12, "FFFFFF", bold=True, italic=True)
    plain_text(slide, 1211.9, 694.2, 71.2, 26.5, str(n), 18, "404040", bold=True, align="ctr")


def chrome_cover(slide, mark):
    bg = slide.background.fill
    bg.solid(); bg.fore_color.rgb = RGBColor.from_string("203864")
    for y in (59.1, 663.7):
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(35.5), E(y), E(1209), E(8.25))
        _fill(shp, "FFFFFF", alpha=0.6); shp.line.fill.background(); shp.shadow.inherit = False
    add_picture(slide, mark, 1075.8, 460.9, 138, 137.6, alpha=0.5)


# ---------------------------------------------------------------- build
def main():
    layout = json.load(open(os.path.join(HERE, "build", "layout.json"), encoding="utf-8"))
    texs = [r["m"] for s in layout["slides"] for it in s["items"] if it["kind"] == "text"
            for r in it["runs"] if "m" in r]
    omml = convert_norm(texs)

    prs, _, _ = load_template()
    logo = os.path.join(SLIDES_DIR, "assets", "logo")
    cover_mark = os.path.join(logo, "smwu-mark-white.png")
    slide_mark = os.path.join(logo, "smwu-mark-blue.png")
    blank = prs.slide_layouts[0]
    order = {"rect": 0, "oval": 1, "tri": 1, "img": 2, "line": 3, "text": 4}
    key = lambda it: 3 if it.get("svg") else order[it["kind"]]   # SVG keeps its paint order
    for s in layout["slides"]:
        slide = prs.slides.add_slide(blank)
        if s["cover"]:
            chrome_cover(slide, cover_mark)
        items = sorted(s["items"], key=key)
        for it in items:
            k = it["kind"]
            if k == "rect":
                add_rect(slide, it)
            elif k == "oval":
                add_oval(slide, it)
            elif k == "tri":
                add_tri(slide, it)
            elif k == "line":
                add_polyline(slide, it)
            elif k == "img":
                add_picture(slide, os.path.join(SLIDES_DIR, it["src"]), it["x"], it["y"], it["w"], it["h"])
            elif k == "text":
                add_text(slide, it, omml)
        if not s["cover"]:
            chrome_content(slide, s["index"], s["title"], slide_mark)
    prs.save(OUT)
    print("saved", OUT, len(prs.slides), "slides")


if __name__ == "__main__":
    main()
