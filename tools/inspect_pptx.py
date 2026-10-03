"""Print layouts, placeholders and shapes of a PPTX (to learn the template geometry).

usage: python inspect_pptx.py FILE.pptx [--slides]
"""
import sys

from pptx import Presentation
from pptx.util import Emu

sys.stdout.reconfigure(encoding="utf-8")
EMU_PX = 9525  # 1 px at 96 dpi


def px(v):
    return round(Emu(v) / EMU_PX, 1) if v is not None else None


prs = Presentation(sys.argv[1])
print("slide size px:", px(prs.slide_width), px(prs.slide_height))
for mi, master in enumerate(prs.slide_masters):
    print(f"MASTER {mi}")
    for sh in master.shapes:
        print("   m-shape", sh.shape_type, sh.name, px(sh.left), px(sh.top), px(sh.width), px(sh.height))
    for li, lay in enumerate(master.slide_layouts):
        print(f"  LAYOUT {li}: {lay.name}")
        for sh in lay.shapes:
            ph = sh.placeholder_format if sh.is_placeholder else None
            tag = f"ph idx={ph.idx} type={ph.type}" if ph else str(sh.shape_type)
            print("     ", tag, "|", sh.name, px(sh.left), px(sh.top), px(sh.width), px(sh.height),
                  "|", (sh.text_frame.text[:40] if sh.has_text_frame else ""))

if "--slides" in sys.argv:
    for si, slide in enumerate(prs.slides, 1):
        print(f"SLIDE {si} layout={slide.slide_layout.name}")
        for sh in slide.shapes:
            print("     ", sh.shape_type, sh.name, px(sh.left), px(sh.top), px(sh.width), px(sh.height),
                  "|", (sh.text_frame.text[:50].replace("\n", " / ") if sh.has_text_frame else ""))
